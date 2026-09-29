import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from report_generator.styles import init_document, add_page_number_field
from report_generator.preliminary import build_preliminary_pages
from report_generator.chapter1 import build_chapter_1
from report_generator.chapter2 import build_chapter_2
from report_generator.chapter3 import build_chapter_3
from report_generator.chapter4 import build_chapter_4
from report_generator.chapter5 import build_chapter_5
from report_generator.chapter6 import build_chapter_6
from report_generator.bibliography import build_bibliography

def main():
    print("Initializing document...")
    doc = init_document()

    print("Building preliminary pages...")
    build_preliminary_pages(doc)

    print("Transitioning to Chapter Section (Arabic page numbering)...")
    # Add new section for Main Body (Chapters 1 to 6 and Bibliography)
    new_section = doc.add_section()
    new_section.top_margin = Inches(1.0)
    new_section.bottom_margin = Inches(1.0)
    new_section.left_margin = Inches(1.25)
    new_section.right_margin = Inches(1.0)
    new_section.page_width = Inches(8.5)
    new_section.page_height = Inches(11.0)

    # Disconnect header and footer from preliminary section
    new_section.header.is_linked_to_previous = False
    new_section.footer.is_linked_to_previous = False

    # Configure Arabic (decimal) numbering starting at page 1
    sectPr = new_section._sectPr
    pgNumType = OxmlElement('w:pgNumType')
    pgNumType.set(qn('w:fmt'), 'decimal')
    pgNumType.set(qn('w:start'), '1')
    sectPr.append(pgNumType)

    # Configure Footer Page Number for Main Chapters
    footer = new_section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f_run = footer_p.add_run()
    f_run.font.name = 'Times New Roman'
    f_run.font.size = Pt(10)
    add_page_number_field(f_run)

    print("Building Chapter 1: Introduction...")
    build_chapter_1(doc)
    doc.add_page_break()

    print("Building Chapter 2: System Analysis...")
    build_chapter_2(doc)
    doc.add_page_break()

    print("Building Chapter 3: System Design...")
    build_chapter_3(doc)
    doc.add_page_break()

    print("Building Chapter 4: System Implementation...")
    build_chapter_4(doc)
    doc.add_page_break()

    print("Building Chapter 5: System Testing and Results...")
    build_chapter_5(doc)
    doc.add_page_break()

    print("Building Chapter 6: Conclusion and Future Enhancements...")
    build_chapter_6(doc)
    doc.add_page_break()

    print("Building Bibliography...")
    build_bibliography(doc)

    output_path = "RouteIQ_Mini_Project_Report.docx"
    doc.save(output_path)
    print(f"Report successfully generated and saved to: {output_path}")

    # Compute Statistics
    total_words = 0
    total_paragraphs = len(doc.paragraphs)
    total_tables = len(doc.tables)

    for p in doc.paragraphs:
        total_words += len(p.text.split())

    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    total_words += len(p.text.split())

    print("\n================ DOCUMENT STATISTICS ================")
    print(f"Total Words: {total_words}")
    print(f"Total Paragraphs: {total_paragraphs}")
    print(f"Total Tables & Figure Boxes: {total_tables}")
    estimated_pages = total_words / 320 + (total_tables * 0.4)
    print(f"Estimated Page Count (TNR 12pt, 1.5 spacing + tables): ~{estimated_pages:.1f} pages")
    print("=====================================================")

if __name__ == "__main__":
    main()
