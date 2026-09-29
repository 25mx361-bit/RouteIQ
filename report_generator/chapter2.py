import docx
from report_generator.styles import (
    add_chapter_heading, add_main_heading, add_sub_heading,
    add_hint_paragraph, add_body_paragraph, add_bullet_point,
    add_table_custom
)

def build_chapter_2(doc):
    add_chapter_heading(doc, "CHAPTER 2\nSYSTEM ANALYSIS")

    # 2.1 Introduction
    add_main_heading(doc, "2.1 Introduction")
    add_hint_paragraph(doc, 
        "System analysis forms the analytical bedrock of software engineering, establishing operational baselines, assessing user demands, and evaluating environmental constraints. This section introduces the systematic analytical methodology applied to the RouteIQ platform."
    )
    add_body_paragraph(doc, 
        "System analysis establishes functional, behavioral, and operational baselines for the software platform. In urban transportation networks, software requirements must address non-stationary telemetry feeds, spatial coordinate geometries, atmospheric factors, and machine learning inference constraints. This chapter contrasts existing navigation platforms against RouteIQ, presents comprehensive functional and non-functional requirements, outlines hardware and software requirements, and details the project feasibility study."
    )

    # 2.2 Existing System
    add_main_heading(doc, "2.2 Existing System")
    add_hint_paragraph(doc, 
        "Evaluating existing systems identifies baseline industry capabilities and architectural patterns. This section reviews conventional routing platforms currently deployed in modern transportation."
    )
    add_body_paragraph(doc, 
        "Prevailing urban transit navigation systems fall into three primary categories:"
    )
    add_bullet_point(doc, "Static Graph-Based Navigation Systems: ", 
        "Traditional routing engines model road networks as static weighted graphs where edges denote street segments and costs reflect physical metric distance or legal speed limits. Algorithms like Dijkstra or standard A* calculate static shortest paths without awareness of live traffic bottlenecks or dynamic delays."
    )
    add_bullet_point(doc, "Instantaneous Real-Time GPS Platforms: ", 
        "Modern commercial platforms (e.g., Google Maps, Apple Maps) leverage crowd-sourced GPS traces and cellular telemetry to compute live segment velocities. While reactive routing reflects current conditions, it assumes that conditions observed at departure will persist throughout the journey, leading to degraded recommendations over longer trips."
    )
    add_bullet_point(doc, "Municipal Intelligent Transportation Systems (ITMS): ", 
        "City traffic management centers deploy inductive loops and camera feeds to synchronize traffic signals. These systems focus primarily on localized signal control rather than multi-attribute route guidance for commuters."
    )

    # 2.3 Limitations of Existing System
    add_main_heading(doc, "2.3 Limitations of Existing System")
    add_hint_paragraph(doc, 
        "Understanding existing deficiencies provides the rationale for architectural innovation. This section enumerates the critical limitations observed in prevailing navigation technologies."
    )
    add_body_paragraph(doc, 
        "Existing navigation systems exhibit several critical limitations:"
    )
    add_bullet_point(doc, "Absence of Long-Horizon Congestion Forecasting: ", 
        "Current platforms cannot project traffic states several hours into the future, relying on historical weekly averages that fail to capture non-stationary trends."
    )
    add_bullet_point(doc, "Susceptibility to Traffic Herding: ", 
        "Diverting large commuter volumes simultaneously onto a single alternate artery creates secondary bottlenecks on narrow corridors."
    )
    add_bullet_point(doc, "Omission of Meteorological Factors: ", 
        "Precipitation, wind gusts, and visibility impairments directly reduce vehicular speed but are omitted from mathematical path scoring."
    )
    add_bullet_point(doc, "Single-Objective Optimization: ", 
        "Routing engines optimize strictly for distance or instantaneous duration, ignoring route reliability and atmospheric risk."
    )

    # 2.4 Proposed System
    add_main_heading(doc, "2.4 Proposed System")
    add_hint_paragraph(doc, 
        "Addressing reactive navigation limitations requires an integrated predictive routing architecture. This section outlines the functional design of the proposed RouteIQ platform."
    )
    add_body_paragraph(doc, 
        "The proposed RouteIQ platform establishes a predictive, multi-criteria route recommendation architecture. RouteIQ introduces a two-stage artificial intelligence pipeline that separates time-series corridor forecasting from multi-attribute route evaluation. Commuter queries trigger automated geocoding via OpenStreetMap and path generation via OpenRouteService. A trained Patch Time Series Transformer (PatchTST) forecasts corridor speeds and congestion indices across a 9-hour lookahead window (36 steps at 15-minute intervals). Subsequently, a CatBoost decision model synthesizes predicted congestion, physical distance, real-time Open-Meteo weather parameters, and administrative disruption notices into an objective route score, recommending the most reliable corridor."
    )

    # 2.5 Advantages of Proposed System
    add_main_heading(doc, "2.5 Advantages of Proposed System")
    add_hint_paragraph(doc, 
        "The proposed architecture provides distinct operational and computational benefits. This section highlights the key advantages of RouteIQ."
    )
    add_bullet_point(doc, "Proactive Long-Horizon Forecasting: ", 
        "Enables commuters to plan future trips (e.g., forecasting 5:00 PM peak conditions from 8:00 AM data) with realistic transit duration estimates."
    )
    add_bullet_point(doc, "High-Capacity PatchTST Modeling: ", 
        "Sub-series patching reduces transformer computational complexity while capturing localized and long-range temporal trends."
    )
    add_bullet_point(doc, "Multi-Attribute CatBoost Ranking: ", 
        "Combines numeric telemetry, categorical weather risk levels, and incident notices without target leakage."
    )
    add_bullet_point(doc, "Integrated Weather Awareness: ", 
        "Open-Meteo atmospheric variables are directly factored into both prediction and route scoring routines."
    )
    add_bullet_point(doc, "Administrative Disruption Governance: ", 
        "Allows municipal authorities to log verified road closures and hazards that immediately update route recommendations."
    )

    # 2.6 Functional Requirements
    add_main_heading(doc, "2.6 Functional Requirements")
    add_hint_paragraph(doc, 
        "Functional requirements define the core software capabilities, inputs, processing rules, and outputs. This section details the twenty primary functional requirements of RouteIQ."
    )
    add_body_paragraph(doc, 
        "The functional capabilities of RouteIQ span user authentication, spatial routing, telemetry harvesting, machine learning forecasting, and administrative governance. These requirements are categorized into six functional modules:"
    )
    add_bullet_point(doc, "Authentication & Session Management: ", 
        "Encompasses FR1 (User Registration) and FR2 (User Login & JWT Session Management) to ensure secure role-based access."
    )
    add_bullet_point(doc, "Geospatial Route Discovery: ", 
        "Encompasses FR3 (Origin Selection), FR4 (Destination Selection), FR5 (Departure Date/Time Configuration), FR6 (OSM Location Geocoding), and FR7 (ORS Multi-Route Generation)."
    )
    add_bullet_point(doc, "Telemetry & Context Ingestion: ", 
        "Encompasses FR8 (Traffic Telemetry Harvesting), FR9 (Open-Meteo Weather Ingestion), and FR10 (Historical Telemetry Processing & 96-step Window Construction)."
    )
    add_bullet_point(doc, "Deep Congestion Forecasting: ", 
        "Encompasses FR11 (PatchTST 9-Hour Congestion Prediction) and FR15 (Accurate Travel Duration Estimation)."
    )
    add_bullet_point(doc, "Route Evaluation & Presentation: ", 
        "Encompasses FR12 (Alternative Route Extraction), FR13 (CatBoost Multi-Criteria Evaluation), FR14 (Recommended Route Badge Display), FR16 (Comparative Matrix Breakdown), FR17 (Interactive Map Visualization), and FR18 (User Search History Logging)."
    )
    add_bullet_point(doc, "Incident Governance: ", 
        "Encompasses FR19 (Administrative Road Disruption Logging) and FR20 (Dynamic Corridor Blocking & Penalty Application)."
    )

    fr_table_data = [
        ["FR1", "User Registration", "Commuter submits username, email, and password.", "User record created; password hashed with bcrypt."],
        ["FR2", "User Login", "Commuter supplies authentication credentials.", "JWT token issued; session state authorized."],
        ["FR3", "Origin Selection", "User selects origin point via text or map pin.", "Origin address string captured."],
        ["FR4", "Destination Selection", "User enters target destination address.", "Destination address string captured."],
        ["FR5", "Date/Time Selection", "User specifies target departure timestamp.", "Timestamp validated and formatted."],
        ["FR6", "Geocoding", "Address strings dispatched to OSM Nominatim.", "WGS84 decimal latitude and longitude returned."],
        ["FR7", "Route Generation", "Coordinates passed to OpenRouteService API.", "Candidate driving polylines and summaries returned."],
        ["FR8", "Traffic Ingestion", "Scheduled daemon queries TomTom / Traffic API.", "Speed, free-flow time, and TCI recorded in DB."],
        ["FR9", "Weather Ingestion", "Corridor coordinates queried via Open-Meteo.", "Rainfall, temperature, wind, humidity recorded."],
        ["FR10", "Historical Processing", "Telemetry pipeline aggregates sensor records.", "96-step chronological context windows assembled."],
        ["FR11", "Congestion Prediction", "Context windows fed into PatchTST PyTorch model.", "36-step future speeds and TCI forecasted."],
        ["FR12", "Alternative Routes", "Routing engine identifies diverse paths.", "Distinct alternative trajectories extracted."],
        ["FR13", "Route Evaluation", "Feature vectors evaluated by CatBoost model.", "Scalar recommendation scores computed."],
        ["FR14", "Recommended Route", "Highest-scoring route highlighted on UI.", "Optimal path tagged with recommendation badge."],
        ["FR15", "Travel Time Display", "Predicted speeds integrated over route distance.", "Accurate forecasted transit duration shown."],
        ["FR16", "Route Comparison", "Comparative metrics tabulated for alternatives.", "Side-by-side trade-off matrix displayed."],
        ["FR17", "Map Visualization", "Polylines drawn on Leaflet/OSM canvas.", "Interactive map with color-coded traffic states."],
        ["FR18", "Search History", "User trip history saved to database.", "Past searches viewable in commuter dashboard."],
        ["FR19", "Admin Updates", "Admin submits road disruption or hazard alert.", "Disruption event stored in admin_updates."],
        ["FR20", "Route Restriction", "Routing pipeline checks corridor against hazards.", "Severely blocked corridors flagged or penalized."]
    ]

    add_table_custom(
        doc,
        "Table 2.1 Functional Requirements Specification",
        ["Req ID", "Requirement Name", "Trigger / Input Description", "System Response / Output"],
        fr_table_data,
        col_widths=[0.8, 1.8, 2.3, 2.3]
    )

    # 2.7 Non-Functional Requirements
    add_main_heading(doc, "2.7 Non-Functional Requirements")
    add_hint_paragraph(doc, 
        "Non-functional requirements specify quality standards, performance thresholds, and operational constraints. This section outlines the quality attributes of RouteIQ."
    )
    add_body_paragraph(doc, 
        "RouteIQ satisfies high standards of software quality to guarantee responsiveness, reliability, and security:"
    )
    add_bullet_point(doc, "Response Latency: ", "End-to-end recommendation generated within 3.0 seconds under standard broadband conditions.")
    add_bullet_point(doc, "Availability: ", "Maintains 99.5% operational uptime during standard operating hours.")
    add_bullet_point(doc, "Scalability: ", "Supports 50 concurrent active commuter route queries without degrading latency.")
    add_bullet_point(doc, "Data Integrity: ", "Enforces relational foreign keys, unique constraints, and ACID transactions in MySQL.")
    add_bullet_point(doc, "Security: ", "Implements salted bcrypt password hashing, signed JWT tokens, and parameterized SQL queries.")
    add_bullet_point(doc, "Usability: ", "Responsive layout across mobile and desktop devices with clear color-coded indicators.")

    nfr_table_data = [
        ["Performance", "End-to-end recommendation response time < 3.0 seconds."],
        ["Availability", "99.5% scheduled operational uptime."],
        ["Scalability", "Support 50 concurrent active commuter queries."],
        ["Reliability", "Graceful fallback if third-party APIs experience temporary rate limiting."],
        ["Security", "Bcrypt password hashing, JWT authorization, sanitized SQL parameters."],
        ["Usability", "Responsive layout across desktop and mobile; clear color-coded indicators."]
    ]
    add_table_custom(
        doc,
        "Table 2.2 Non-Functional Requirements Specification",
        ["Quality Attribute", "Specification and Operational Target"],
        nfr_table_data,
        col_widths=[2.0, 5.2]
    )

    # 2.8 Hardware Requirements
    add_main_heading(doc, "2.8 Hardware Requirements")
    add_hint_paragraph(doc, 
        "Executing deep learning models alongside web persistence requires balanced hardware infrastructure. This section outlines hardware configurations."
    )
    hw_table_data = [
        ["Server / Backend Host", "Intel Core i5 (4 Cores)", "Intel Core i7 (8 Cores)", "2.5 GHz+ base clock"],
        ["System Memory (RAM)", "8 GB DDR4", "16 GB DDR4/DDR5", "Accommodates PyTorch, Flask, MySQL"],
        ["Graphics Processing (GPU)", "Integrated GPU (CPU inference)", "NVIDIA RTX 3060 (6GB+ VRAM)", "CUDA accelerated training for PatchTST"],
        ["Storage Capacity", "20 GB Available SSD Space", "50 GB NVMe M.2 SSD", "High IOPS for telemetry & weights"],
        ["Client Commuter Device", "Dual Core 1.8 GHz, 4 GB RAM", "Modern Smartphone or PC", "Web browser with WebGL canvas support"]
    ]
    add_table_custom(
        doc,
        "Table 2.3 Hardware Requirements for Server, Training, and Client",
        ["Component", "Minimum Specification", "Recommended Specification", "Operational Purpose"],
        hw_table_data,
        col_widths=[1.8, 1.8, 1.8, 1.8]
    )

    # 2.9 Software Requirements
    add_main_heading(doc, "2.9 Software Requirements")
    add_hint_paragraph(doc, 
        "Stable execution requires compatible operating systems, runtimes, database servers, and framework versions. This section details the software stack."
    )
    sw_table_data = [
        ["Operating System", "Microsoft Windows 10/11 (64-bit) / Ubuntu Linux 22.04 LTS"],
        ["Backend Runtime", "Python 3.10+ (64-bit)"],
        ["Web Framework", "Flask v2.3+, Flask-CORS, Flask-Bcrypt, PyJWT"],
        ["Relational Database", "MySQL Server v8.0+ / PyMySQL connector"],
        ["Deep Learning Engine", "PyTorch v2.0+ (torch, torchvision) with CUDA support"],
        ["Machine Learning Suite", "CatBoost v1.2+, Scikit-Learn v1.3+, Pandas v2.0+, NumPy v1.24+"],
        ["Frontend Environment", "Node.js v18+, React.js v18.2+, Vite build system"],
        ["UI Styling & Icons", "Tailwind CSS v3.3+, PostCSS, Lucide React Icons"],
        ["Mapping Canvas Library", "Leaflet.js v1.9+, React-Leaflet"]
    ]
    add_table_custom(
        doc,
        "Table 2.4 Software Stack and Library Dependencies",
        ["Software Component / Tool", "Version and Technical Specification"],
        sw_table_data,
        col_widths=[2.5, 4.7]
    )

    # 2.10 Feasibility Study
    add_main_heading(doc, "2.10 Feasibility Study")
    add_hint_paragraph(doc, 
        "A formal feasibility analysis validates project completion within technical, financial, and operational constraints. This section examines key feasibility dimensions."
    )
    add_body_paragraph(doc, 
        "The feasibility assessment confirmed that RouteIQ is technically, economically, and operationally viable:"
    )
    add_bullet_point(doc, "Technical Feasibility: ", "The algorithms (PatchTST and CatBoost) are supported by mature open-source implementations. OpenStreetMap, OpenRouteService, and Open-Meteo APIs provide verified spatial and weather feeds without custom satellites.")
    add_bullet_point(doc, "Economic Feasibility: ", "The project relies entirely on open-source frameworks and public API tiers, requiring no proprietary software licenses.")
    add_bullet_point(doc, "Operational Feasibility: ", "The web interface provides an intuitive search flow suitable for all commuter skill levels, while administrative portals use structured forms.")
    add_bullet_point(doc, "Schedule Feasibility: ", "The modular architecture allowed sequential delivery of database, routing, ML forecasting, and frontend tiers within the academic timeframe.")

    # 2.11 User Roles
    add_main_heading(doc, "2.11 User Roles")
    add_hint_paragraph(doc, 
        "Role-based access separates commuter activities from administrative management. This section defines the privileges allocated to system actors."
    )
    add_bullet_point(doc, "Commuter / General User: ", "Can register accounts, authenticate sessions, search routes, inspect forecasted congestion levels, evaluate comparative metrics, and review query history.")
    add_bullet_point(doc, "System Administrator: ", "Possesses elevated privileges to log road closures, monitor telemetry collection health, trigger model retraining, and inspect validation metrics.")

    # 2.12 System Constraints
    add_main_heading(doc, "2.12 System Constraints")
    add_hint_paragraph(doc, 
        "Engineering designs must accommodate realistic technological and real-world boundaries. This section outlines system constraints."
    )
    add_bullet_point(doc, "API Rate Limits: ", "Public geocoding endpoints require throttled requests, managed via in-memory caching.")
    add_bullet_point(doc, "Network Dependency: ", "Real-time atmospheric and routing updates require active internet connectivity.")
    add_bullet_point(doc, "Geographical Scope: ", "Empirical validation was conducted on monitored corridors in Coimbatore, Tamil Nadu.")

    # 2.13 Concluding Remarks
    add_main_heading(doc, "2.13 Concluding Remarks")
    add_hint_paragraph(doc, 
        "This section concludes the system analysis chapter, summarizing analytical specifications and transitioning to architectural design."
    )
    add_body_paragraph(doc, 
        "This chapter established the system analysis for RouteIQ. The deficiencies of reactive navigation systems were analyzed, motivating a forward-looking predictive framework. Twenty functional requirements, quality attributes, hardware/software specifications, and feasibility dimensions were detailed, laying the groundwork for the architectural design presented in Chapter 3."
    )
