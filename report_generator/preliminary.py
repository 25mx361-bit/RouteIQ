import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from report_generator.styles import (
    add_chapter_heading, add_body_paragraph, add_table_custom
)

def build_preliminary_pages(doc):
    # ── COVER PAGE ────────────────────────────────────────────────────────────
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(20)
    run_t = p_title.add_run("ROUTEIQ: SMART TRAFFIC CONGESTION PREDICTION AND ROUTE RECOMMENDATION SYSTEM\n")
    run_t.font.name = 'Times New Roman'
    run_t.font.size = Pt(16)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(10)
    p_sub.paragraph_format.space_after = Pt(24)
    run_sub = p_sub.add_run("A MINI PROJECT REPORT\n\nSubmitted by\n\n")
    run_sub.font.name = 'Times New Roman'
    run_sub.font.size = Pt(12)
    
    run_names = p_sub.add_run("SANJANA M (25MX361)\nSRIVIKASHNI S (25MX350)\nTEAM 33\n\n")
    run_names.font.name = 'Times New Roman'
    run_names.font.size = Pt(13)
    run_names.font.bold = True

    run_deg = p_sub.add_run(
        "in partial fulfillment for the award of the degree\nof\n"
        "MASTER OF COMPUTER APPLICATIONS\n"
        "OF ANNA UNIVERSITY\n"
    )
    run_deg.font.name = 'Times New Roman'
    run_deg.font.size = Pt(12)

    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(36)
    p_inst.paragraph_format.space_after = Pt(12)
    run_inst = p_inst.add_run(
        "DEPARTMENT OF COMPUTER APPLICATIONS\n"
        "PSG COLLEGE OF TECHNOLOGY\n"
        "(Autonomous Institution)\n"
        "COIMBATORE – 641 004\n\n"
        "ACADEMIC YEAR: 2026–2027\n"
        "OCTOBER 2026"
    )
    run_inst.font.name = 'Times New Roman'
    run_inst.font.size = Pt(12)
    run_inst.font.bold = True

    doc.add_page_break()

    # ── BONAFIDE CERTIFICATE ──────────────────────────────────────────────────
    p_cert_inst = doc.add_paragraph()
    p_cert_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cert_inst.paragraph_format.space_before = Pt(18)
    p_cert_inst.paragraph_format.space_after = Pt(12)
    run_ci = p_cert_inst.add_run(
        "PSG COLLEGE OF TECHNOLOGY\n"
        "(Autonomous Institution)\n"
        "COIMBATORE – 641 004\n\n"
        "DEPARTMENT OF COMPUTER APPLICATIONS\n"
        "BONAFIDE CERTIFICATE"
    )
    run_ci.font.name = 'Times New Roman'
    run_ci.font.size = Pt(14)
    run_ci.font.bold = True

    p_cert_body = doc.add_paragraph()
    p_cert_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_cert_body.paragraph_format.line_spacing = 1.5
    p_cert_body.paragraph_format.space_before = Pt(16)
    p_cert_body.paragraph_format.space_after = Pt(14)
    run_cb = p_cert_body.add_run(
        "This is to certify that the project work entitled \"ROUTEIQ: SMART TRAFFIC CONGESTION PREDICTION AND ROUTE RECOMMENDATION SYSTEM\" "
        "is a bonafide record of work done by SANJANA M (Roll No: 25MX361) and SRIVIKASHNI S (Roll No: 25MX350) of Team 33 in "
        "partial fulfillment of the requirements for the award of the degree of MASTER OF COMPUTER APPLICATIONS of Anna University, "
        "Chennai, carried out at PSG College of Technology, Coimbatore during the academic year 2026–2027."
    )
    run_cb.font.name = 'Times New Roman'
    run_cb.font.size = Pt(12)

    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_sig.paragraph_format.space_before = Pt(40)
    p_sig.paragraph_format.space_after = Pt(12)
    run_sig = p_sig.add_run(
        "FACULTY GUIDE                                                  HEAD OF THE DEPARTMENT\n\n"
        "[INSERT GUIDE NAME]                                          Dr. A. CHITRA\n"
        "Assistant Professor                                           Professor and Head\n"
        "Department of Computer Applications                           Department of Computer Applications\n"
        "PSG College of Technology                                     PSG College of Technology\n"
        "Coimbatore – 641 004                                          Coimbatore – 641 004\n\n\n"
        "Submitted for the Mini Project Viva-Voce examination held on ____________________\n\n\n"
        "INTERNAL EXAMINER                                              EXTERNAL EXAMINER"
    )
    run_sig.font.name = 'Times New Roman'
    run_sig.font.size = Pt(11)
    run_sig.font.bold = True

    doc.add_page_break()

    # ── ACKNOWLEDGEMENT ───────────────────────────────────────────────────────
    add_chapter_heading(doc, "ACKNOWLEDGEMENT")
    
    add_body_paragraph(doc, 
        "The project members take this opportunity to express profound gratitude to the esteemed Principal, "
        "Dr. K. PRAKASAN, PSG College of Technology, for providing the necessary institutional facilities, academic infrastructure, "
        "and encouraging atmosphere required for carrying out this mini project successfully."
    )
    add_body_paragraph(doc, 
        "Sincere and deep gratitude is expressed to Dr. A. CHITRA, Professor and Head, Department of Computer Applications, "
        "PSG College of Technology, for constant motivation, visionary academic guidance, and invaluable support throughout the project."
    )
    add_body_paragraph(doc, 
        "Warm thanks are extended to the Programme Coordinator, Dr. R. MANAVALAN, Associate Professor, Department of Computer Applications, "
        "PSG College of Technology, whose continuous academic support, guidance, and evaluation encouraged the timely execution of the project milestones."
    )
    add_body_paragraph(doc, 
        "Heartfelt gratitude is rendered to the Faculty Guide, [INSERT GUIDE NAME], Assistant Professor, Department of Computer Applications, "
        "PSG College of Technology, for providing continuous technical direction, insightful feedback, constructive reviews, and enduring encouragement during the entire development life cycle."
    )
    add_body_paragraph(doc, 
        "Gratitude is also extended to the Tutor, Dr. M. SUBATHRA, Assistant Professor (Sr. Gr.), Department of Computer Applications, "
        "PSG College of Technology, for her persistent guidance, monitoring, and helpful advice throughout the tenure of the mini project."
    )
    add_body_paragraph(doc, 
        "The project team conveys sincere thanks to all the faculty members and technical staff of the Department of Computer Applications "
        "for their kind cooperation, intellectual support, and guidance. Finally, warmest appreciation is conveyed to parents, family members, "
        "and friends whose enduring encouragement, sacrifices, and belief served as a pillar of strength throughout this academic endeavor."
    )

    doc.add_page_break()

    # ── SYNOPSIS ──────────────────────────────────────────────────────────────
    add_chapter_heading(doc, "SYNOPSIS")
    
    add_body_paragraph(doc, 
        "Rapid urbanization, expanding vehicle ownership, and unpredictable road events have escalated urban traffic congestion into a "
        "critical societal and environmental challenge. Traditional vehicular navigation systems primarily rely on static graph-based shortest "
        "path algorithms or instantaneous real-time traffic speeds, often directing vehicular flows onto corridors that become severely bottlenecked "
        "shortly after user departure. To mitigate these pervasive challenges, this project presents RouteIQ, an intelligent, multi-modal web-based "
        "traffic congestion prediction and route recommendation system developed for modern urban mobility management."
    )
    add_body_paragraph(doc, 
        "The proposed system combines deep-learning time-series forecasting with gradient-boosted decision trees to evaluate candidate urban routes "
        "holistically. RouteIQ utilizes OpenStreetMap (OSM) via Nominatim for spatial geocoding and OpenRouteService (ORS) for generating diverse "
        "candidate trajectories between user-specified origins and destinations. To account for environmental and temporal conditions, the system "
        "integrates meteorological features retrieved in real time through the Open-Meteo API alongside live and historical traffic telemetry."
    )
    add_body_paragraph(doc, 
        "The core intelligence of RouteIQ is governed by a two-stage machine learning pipeline. In the first stage, PatchTST (Patch Time Series Transformer) "
        "is implemented to perform multi-step time-series forecasting. Operating on chronological sequences of traffic speed, travel time, and environmental "
        "variables segmented into distinct temporal patches, PatchTST forecasts road corridor congestion up to nine hours ahead (e.g., forecasting 5:00 PM "
        "traffic from 8:00 AM observations). In the second stage, CatBoost, an advanced gradient-boosted decision tree architecture capable of handling "
        "both numerical and categorical features without target leakage, computes composite multi-criteria route scores. These scores integrate predicted "
        "congestion indices, travel durations, physical road distances, weather hazards, and administrator-reported exceptional road disruptions "
        "(such as construction zones, public assemblies, or landslides)."
    )
    add_body_paragraph(doc, 
        "The system is engineered as a decoupled, responsive web application comprising a high-performance Python Flask REST backend, a normalized MySQL "
        "relational database, and an intuitive user interface built with React.js and Tailwind CSS. RouteIQ empowers urban commuters with forward-looking "
        "congestion awareness, transparent comparative analytics, and proactive route recommendations, paving the way for optimized urban mobility and "
        "reduced transit delays."
    )

    doc.add_page_break()

    # ── TABLE OF CONTENTS ─────────────────────────────────────────────────────
    add_chapter_heading(doc, "TABLE OF CONTENTS")
    
    toc_data = [
        ("ACKNOWLEDGEMENT", "i"),
        ("SYNOPSIS", "ii"),
        ("LIST OF FIGURES", "iv"),
        ("LIST OF TABLES", "v"),
        ("LIST OF ABBREVIATIONS", "vi"),
        ("CHAPTER 1: INTRODUCTION", "1"),
        ("    1.1 Project Overview", "1"),
        ("    1.2 Background", "2"),
        ("    1.3 Problem Statement", "2"),
        ("    1.4 Need for the Project", "3"),
        ("    1.5 Objectives", "3"),
        ("    1.6 Scope of the Project", "4"),
        ("    1.7 Significance of the Project", "4"),
        ("    1.8 Project Workflow", "5"),
        ("    1.9 Technologies Used", "5"),
        ("    1.10 Concluding Remarks", "6"),
        ("CHAPTER 2: SYSTEM ANALYSIS", "7"),
        ("    2.1 Introduction", "7"),
        ("    2.2 Existing System", "7"),
        ("    2.3 Limitations of Existing System", "8"),
        ("    2.4 Proposed System", "8"),
        ("    2.5 Advantages of Proposed System", "9"),
        ("    2.6 Functional Requirements", "9"),
        ("    2.7 Non-Functional Requirements", "11"),
        ("    2.8 Hardware Requirements", "12"),
        ("    2.9 Software Requirements", "12"),
        ("    2.10 Feasibility Study", "13"),
        ("    2.11 User Roles", "13"),
        ("    2.12 System Constraints", "14"),
        ("    2.13 Concluding Remarks", "14"),
        ("CHAPTER 3: SYSTEM DESIGN", "15"),
        ("    3.1 Introduction", "15"),
        ("    3.2 Overall System Architecture", "15"),
        ("    3.3 System Workflow", "16"),
        ("    3.4 Use Case Diagram", "17"),
        ("    3.5 Activity Diagram", "18"),
        ("    3.6 Data Flow Diagram", "18"),
        ("    3.7 Sequence Diagram", "19"),
        ("    3.8 Component Diagram", "20"),
        ("    3.9 Deployment Diagram", "20"),
        ("    3.10 Database Design", "21"),
        ("    3.11 ER Diagram", "21"),
        ("    3.12 Database Tables", "22"),
        ("    3.13 API Architecture", "23"),
        ("    3.14 Machine Learning Pipeline", "24"),
        ("    3.15 PatchTST Architecture", "25"),
        ("    3.16 CatBoost Route Recommendation Architecture", "26"),
        ("    3.17 Concluding Remarks", "26"),
        ("CHAPTER 4: SYSTEM IMPLEMENTATION", "27"),
        ("    4.1 Introduction", "27"),
        ("    4.2 Development Environment", "27"),
        ("    4.3 Frontend Implementation", "28"),
        ("    4.4 Backend Implementation", "28"),
        ("    4.5 Database Implementation", "29"),
        ("    4.6 User Registration and Login", "29"),
        ("    4.7 Source and Destination Module", "30"),
        ("    4.8 Geocoding Module", "30"),
        ("    4.9 Route Generation Module", "31"),
        ("    4.10 Traffic Data Collection", "31"),
        ("    4.11 Weather Data Collection", "32"),
        ("    4.12 Historical Dataset Processing", "32"),
        ("    4.13 Data Preprocessing", "33"),
        ("    4.14 Feature Engineering", "33"),
        ("    4.15 PatchTST Implementation", "34"),
        ("    4.16 Future Traffic Prediction", "34"),
        ("    4.17 CatBoost Implementation", "35"),
        ("    4.18 Route Recommendation", "35"),
        ("    4.19 Alternative Route Comparison", "36"),
        ("    4.20 Dashboard", "36"),
        ("    4.21 Admin Module", "37"),
        ("    4.22 Exceptional Road Condition Handling", "37"),
        ("    4.23 API Integration", "38"),
        ("    4.24 Database Operations", "38"),
        ("    4.25 Security Considerations", "39"),
        ("    4.26 Important Code Snippets", "39"),
        ("    4.27 Concluding Remarks", "40"),
        ("CHAPTER 5: SYSTEM TESTING AND RESULTS", "41"),
        ("    5.1 Introduction", "41"),
        ("    5.2 Testing Strategy", "41"),
        ("    5.3 Unit Testing", "42"),
        ("    5.4 Integration Testing", "42"),
        ("    5.5 System Testing", "43"),
        ("    5.6 Functional Testing", "43"),
        ("    5.7 API Testing", "43"),
        ("    5.8 Database Testing", "44"),
        ("    5.9 Machine Learning Model Testing", "44"),
        ("    5.10 Prediction Validation", "45"),
        ("    5.11 Route Recommendation Testing", "45"),
        ("    5.12 Performance Testing", "45"),
        ("    5.13 Usability Testing", "46"),
        ("    5.14 Test Cases", "46"),
        ("    5.15 Test Results", "47"),
        ("    5.16 Error Handling", "48"),
        ("    5.17 Limitations", "48"),
        ("    5.18 Concluding Remarks", "48"),
        ("CHAPTER 6: CONCLUSION AND FUTURE ENHANCEMENTS", "49"),
        ("    6.1 Introduction", "49"),
        ("    6.2 Summary of the Work", "49"),
        ("    6.3 Major Contributions", "50"),
        ("    6.4 System Outcomes", "50"),
        ("    6.5 Limitations", "51"),
        ("    6.6 Future Enhancements", "51"),
        ("    6.7 Conclusion", "52"),
        ("    6.8 Final Concluding Remarks", "52"),
        ("BIBLIOGRAPHY", "53")
    ]

    add_table_custom(
        doc,
        "Table of Contents Hierarchy",
        ["Topic / Section Heading", "Page"],
        [[item[0], item[1]] for item in toc_data],
        col_widths=[6.0, 1.2]
    )

    doc.add_page_break()

    # ── LIST OF FIGURES ───────────────────────────────────────────────────────
    add_chapter_heading(doc, "LIST OF FIGURES")
    
    figures_list = [
        ("Figure 1.1", "High-Level Architectural Flow of RouteIQ", "5"),
        ("Figure 3.1", "Overall System Architecture of RouteIQ", "15"),
        ("Figure 3.2", "System Workflow Diagram from Input to Route Recommendation", "16"),
        ("Figure 3.3", "Use Case Diagram for RouteIQ Actors", "17"),
        ("Figure 3.4", "Activity Diagram of Commuter Route Search and Evaluation", "18"),
        ("Figure 3.5", "Context Analysis: Data Flow Diagram (Level 0)", "18"),
        ("Figure 3.6", "Detailed Functional Flow: Data Flow Diagram (Level 1)", "19"),
        ("Figure 3.7", "Sequence Diagram for Predictive Route Recommendation Request", "19"),
        ("Figure 3.8", "Component Diagram of RouteIQ System Modules", "20"),
        ("Figure 3.9", "Deployment Diagram of RouteIQ Infrastructure", "20"),
        ("Figure 3.10", "Entity Relationship (ER) Diagram of Relational Database Schema", "21"),
        ("Figure 3.11", "API Interaction and Gateway Architecture", "23"),
        ("Figure 3.12", "End-to-End Machine Learning Pipeline of RouteIQ", "24"),
        ("Figure 3.13", "PatchTST Patch-Based Transformer Architecture for Time-Series Traffic Forecasting", "25"),
        ("Figure 3.14", "CatBoost Decision Forest for Multi-Attribute Route Recommendation", "26"),
        ("Figure 4.1", "Home and Navigation Landing Interface [INSERT SCREENSHOT HERE]", "28"),
        ("Figure 4.2", "User Authentication and Registration Interface [INSERT SCREENSHOT HERE]", "29"),
        ("Figure 4.3", "Interactive Route Search and Geocoding View [INSERT SCREENSHOT HERE]", "30"),
        ("Figure 4.4", "Candidate Alternative Route Display [INSERT SCREENSHOT HERE]", "31"),
        ("Figure 4.5", "Weather and Environmental Impact Panel [INSERT SCREENSHOT HERE]", "32"),
        ("Figure 4.6", "PatchTST 9-Hour Congestion Forecasting Visualizer [INSERT SCREENSHOT HERE]", "34"),
        ("Figure 4.7", "CatBoost Comparative Route Ranking and Recommendation [INSERT SCREENSHOT HERE]", "35"),
        ("Figure 4.8", "Administrative Incident and Road Disruption Control Panel [INSERT SCREENSHOT HERE]", "37")
    ]

    add_table_custom(
        doc,
        "List of Figures",
        ["Figure Number", "Figure Title", "Page"],
        [[f[0], f[1], f[2]] for f in figures_list],
        col_widths=[1.5, 5.0, 0.8]
    )

    doc.add_page_break()

    # ── LIST OF TABLES ────────────────────────────────────────────────────────
    add_chapter_heading(doc, "LIST OF TABLES")
    
    tables_list = [
        ("Table 2.1", "Functional Requirements Specification", "10"),
        ("Table 2.2", "Non-Functional Requirements Specification", "11"),
        ("Table 2.3", "Hardware Requirements for Server, Training, and Client", "12"),
        ("Table 2.4", "Software Stack and Library Dependencies", "12"),
        ("Table 3.1", "Database Table: users", "22"),
        ("Table 3.2", "Database Table: search_history", "22"),
        ("Table 3.3", "Database Table: routes", "22"),
        ("Table 3.4", "Database Table: traffic_data", "22"),
        ("Table 3.5", "Database Table: weather_data", "23"),
        ("Table 3.6", "Database Table: predictions", "23"),
        ("Table 3.7", "Database Table: route_recommendations", "23"),
        ("Table 3.8", "Database Table: admin_updates", "23"),
        ("Table 4.1", "External API Specifications and Operational Roles", "38"),
        ("Table 5.1", "Comprehensive System Test Cases and Functional Validation", "46"),
        ("Table 5.2", "PatchTST Traffic Prediction Quantitative Performance Metrics", "47"),
        ("Table 5.3", "Baseline Comparative Performance for Corridor Speed Forecasting", "47")
    ]

    add_table_custom(
        doc,
        "List of Tables",
        ["Table Number", "Table Title", "Page"],
        [[t[0], t[1], t[2]] for t in tables_list],
        col_widths=[1.5, 5.0, 0.8]
    )

    doc.add_page_break()

    # ── LIST OF ABBREVIATIONS ─────────────────────────────────────────────────
    add_chapter_heading(doc, "LIST OF ABBREVIATIONS")
    
    abbrev_list = [
        ("API", "Application Programming Interface"),
        ("CatBoost", "Categorical Boosting"),
        ("DFD", "Data Flow Diagram"),
        ("ER", "Entity Relationship"),
        ("FFN", "Feed-Forward Network"),
        ("GNN", "Graph Neural Network"),
        ("GPS", "Global Positioning System"),
        ("HTTP", "Hypertext Transfer Protocol"),
        ("JSON", "JavaScript Object Notation"),
        ("JWT", "JSON Web Token"),
        ("MAE", "Mean Absolute Error"),
        ("MAPE", "Mean Absolute Percentage Error"),
        ("ML", "Machine Learning"),
        ("ORS", "OpenRouteService"),
        ("OSM", "OpenStreetMap"),
        ("PatchTST", "Patch Time Series Transformer"),
        ("REST", "Representational State Transfer"),
        ("RMSE", "Root Mean Squared Error"),
        ("SQL", "Structured Query Language"),
        ("TCI", "Traffic Congestion Index"),
        ("TNR", "Times New Roman"),
        ("UI", "User Interface"),
        ("UML", "Unified Modeling Language"),
        ("WGS84", "World Geodetic System 1984")
    ]

    add_table_custom(
        doc,
        "List of Abbreviations",
        ["Abbreviation", "Expanded Academic Definition"],
        [[a[0], a[1]] for a in abbrev_list],
        col_widths=[2.0, 5.2]
    )
